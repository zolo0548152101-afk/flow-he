"""ייצוא תמלול לקובץ Word בעברית: פסקאות מימין לשמאל (bidi + rtl), דובר מודגש ושעה."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "Arial"


def _rtl_par(p):
    # בלי jc: בפסקה bidi ברירת המחדל היא "התחלה" = ימין. "right" מפורש מתהפך לשמאל ב-LibreOffice
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1"); pPr.append(bidi)


def _run(p, text, bold=False, color=None, size=12):
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts"); rPr.append(fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(k), FONT)
    rtl = OxmlElement("w:rtl"); rtl.set(qn("w:val"), "1"); rPr.append(rtl)
    if bold:
        bcs = OxmlElement("w:bCs"); rPr.append(bcs)
    szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(size * 2)); rPr.append(szcs)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def fmt_time(sec):
    sec = int(sec)
    return f"{sec // 3600}:{sec // 60 % 60:02d}:{sec % 60:02d}" if sec >= 3600 else f"{sec // 60}:{sec % 60:02d}"


SPEAKER_COLORS = [(91, 70, 230), (14, 116, 144), (190, 70, 40), (40, 130, 60)]


def write_docx(path, title, subtitle, entries, show_speakers=True):
    """entries: [(start_sec, speaker, text)] לפי הסדר."""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(12)
    h = doc.add_paragraph(); _rtl_par(h); _run(h, title, bold=True, size=18)
    if subtitle:
        s = doc.add_paragraph(); _rtl_par(s); _run(s, subtitle, color=(120, 120, 130), size=10)
    colors = {}
    prev = None
    for start, spk, text in entries:
        if not text:
            continue
        p = doc.add_paragraph(); _rtl_par(p)
        p.paragraph_format.space_after = Pt(6)
        if show_speakers and spk != prev:
            col = colors.setdefault(spk, SPEAKER_COLORS[len(colors) % len(SPEAKER_COLORS)])
            _run(p, f"{spk}  ", bold=True, color=col)
            _run(p, f"{fmt_time(start)}  ", color=(150, 150, 160), size=9)
            p.add_run().add_break()
        prev = spk
        _run(p, text)
    doc.save(path)
    return path
