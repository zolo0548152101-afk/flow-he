"""סיכום הקלטה (פגישה / שיעור / שיחה) עם Claude.

המשתמש בוחר שני דברים בלבד: סוג השיחה ורמת פירוט. לכל סוג יש מבנה סיכום משלו.
"""
import re

TYPES = [
    ("work", "פגישת עבודה", "החלטות, משימות ומי אחראי",
     "## בשורה אחת\nמה הייתה מטרת הפגישה ומה יצא ממנה.\n"
     "## החלטות\n"
     "## משימות\nכל משימה בשורה: מי · מה · עד מתי (רק אם נאמר).\n"
     "## מה נדון\nהנושאים לפי הסדר.\n"
     "## פתוח / לבדוק\nשאלות שנשארו בלי תשובה."),
    ("lecture", "הרצאה / שיעור", "רעיונות, מושגים ושאלות לחזרה",
     "## הנושא בשורה אחת\n"
     "## רעיונות מרכזיים\n"
     "## מושגים והגדרות\nכל מושג: הסבר קצר במילים פשוטות.\n"
     "## דוגמאות שהוזכרו\n"
     "## שאלות לחזרה\n3–5 שאלות שבודקות הבנה של החומר."),
    ("call", "שיחה אישית / טלפון", "על מה דיברו ומה סוכם",
     "## על מה דיברנו\n"
     "## מה סוכם\nתוכניות, הבטחות, מפגשים.\n"
     "## לזכור\nתאריכים, שמות, מספרים, המלצות שעלו."),
    ("client", "שיחת לקוח / ראיון", "צרכים, ציטוטים והמשך",
     "## מי ומה\nמי הלקוח/המרואיין ומה הרקע.\n"
     "## צרכים ובעיות\n"
     "## ציטוטים חשובים\nמשפטים מדויקים מהתמלול, במירכאות.\n"
     "## צעדים הבאים"),
    ("ideas", "סיעור מוחות", "רעיונות שעלו ומה נבחר",
     "## השאלה\nעל מה חשבו יחד.\n"
     "## הרעיונות שעלו\nמקובצים לפי נושא.\n"
     "## הכיוון שנבחר\n"
     "## מה לבדוק עכשיו"),
]

LENGTHS = [
    ("short", "קצר", "רק העיקר. עד כ-150 מילים, נקודות קצרות. השמט סעיף שאין בו משהו חשוב."),
    ("detailed", "מפורט", "כסה את כל מה שנאמר לפי הסדר, עם פרטים, שמות ומספרים. אפשר תתי-נקודות (שתי רווחים ואז '- ')."),
]

SYSTEM = (
    "אתה מסכם תמלולים בעברית. התמלול נוצר אוטומטית מדיבור ועלול להכיל שגיאות זיהוי: "
    "אם מילה לא הגיונית — הבן מההקשר ואל תצטט אותה. אל תמציא דבר שלא נאמר. "
    "הדוברים מסומנים בתוויות כמו 'דובר 1' או 'אני'; אם בשיחה עצמה נאמר שם של מישהו — השתמש בשם. "
    "כתוב עברית ברורה ותקנית.\n"
    "פורמט: כותרות בשורה שמתחילה ב-'## ', נקודות בשורה שמתחילה ב-'- '. "
    "בלי טבלאות, בלי ** ובלי הקדמה או סיום משלך. אם לסעיף אין תוכן — השמט אותו."
)


def type_by_key(k):
    return next((t for t in TYPES if t[0] == k), TYPES[0])


def length_by_key(k):
    return next((x for x in LENGTHS if x[0] == k), LENGTHS[0])


def build_prompt(kind, length, transcript, title=""):
    t = type_by_key(kind)
    ln = length_by_key(length)
    user = (f"סוג השיחה: {t[1]}.\nרמת פירוט: {ln[1]} — {ln[2]}\n\n"
            f"מבנה הסיכום:\n{t[3]}\n\n"
            + (f"כותרת ההקלטה: {title}\n\n" if title else "")
            + "התמלול:\n<<<\n" + transcript[-350000:] + "\n>>>")
    return SYSTEM, user


def fmt_time(sec):
    sec = int(sec)
    return f"{sec // 3600}:{sec // 60 % 60:02d}:{sec % 60:02d}" if sec >= 3600 else f"{sec // 60}:{sec % 60:02d}"


def transcript_from_entries(entries):
    """entries: [(start, speaker, text, ...)] -> טקסט עם דובר וזמן בכל מעבר דובר."""
    out, prev = [], None
    for e in entries:
        start, spk, text = e[0], e[1], e[2]
        if not text:
            continue
        if spk != prev:
            out.append(f"\n[{fmt_time(start)}] {spk}:")
            prev = spk
        out.append(text)
    return "\n".join(out).strip()


def transcript_from_docx(path):
    from docx import Document
    return "\n".join(p.text for p in Document(path).paragraphs if p.text.strip())


def parse(md):
    """[(kind, text)] כש-kind הוא h / li / li2 / p."""
    out = []
    for line in (md or "").splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        s = s.replace("**", "")
        if re.match(r"^\s*#{1,6}\s", s):
            out.append(("h", re.sub(r"^\s*#+\s*", "", s)))
        elif re.match(r"^\s{2,}[-•*]\s", s):
            out.append(("li2", re.sub(r"^\s*[-•*]\s*", "", s)))
        elif re.match(r"^\s*[-•*]\s", s):
            out.append(("li", re.sub(r"^\s*[-•*]\s*", "", s)))
        else:
            out.append(("p", s.strip()))
    return out


def write_docx(path, title, subtitle, md):
    import meeting_doc as md_
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = md_.FONT; st.font.size = Pt(12)
    h = doc.add_paragraph(); md_._rtl_par(h); md_._run(h, title, bold=True, size=18)
    if subtitle:
        s = doc.add_paragraph(); md_._rtl_par(s); md_._run(s, subtitle, color=(120, 120, 130), size=10)
    for kind, text in parse(md):
        p = doc.add_paragraph(); md_._rtl_par(p)
        if kind == "h":
            p.paragraph_format.space_before = Pt(12)
            md_._run(p, text, bold=True, size=14, color=(91, 70, 230))
        elif kind in ("li", "li2"):
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            ind = OxmlElement("w:ind"); ind.set(qn("w:start"), "280" if kind == "li" else "640")   # הזחה מצד ההתחלה (ימין)
            p._p.get_or_add_pPr().append(ind)
            p.paragraph_format.space_after = Pt(3)
            md_._run(p, ("• " if kind == "li" else "◦ ") + text)
        else:
            md_._run(p, text)
    doc.save(path)
    return path
